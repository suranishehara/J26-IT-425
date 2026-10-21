"""Non-destructive volunteer intake; permission is user-reported, not inferred from XLSX.

Run with a Python environment containing openpyxl. Training is deliberately separate.
"""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
from difflib import SequenceMatcher
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from backend.nlp import normalize
from scripts.prepare_nlu_data import INTENTS

HEADERS = ('Participant code', 'Message ID', 'Conversation ID', 'Message', 'Intended purpose')


def text_key(text):
    return re.sub(r'\s+', ' ', normalize(text)).strip()


def read_records(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Messages' not in workbook.sheetnames:
            raise ValueError('Expected Messages sheet.')
        sheet = workbook['Messages']
        rows = sheet.iter_rows(values_only=False)
        header = next(rows)
        positions = {str(cell.value).strip(): index for index, cell in enumerate(header)}
        if not set(HEADERS).issubset(positions):
            raise ValueError('Missing required source columns.')
        records = []
        for number, cells in enumerate(rows, 2):
            if not any(cell.value is not None for cell in cells):
                continue
            selected = [cells[positions[name]] for name in HEADERS]
            if any(cell.data_type == 'f' for cell in selected):
                raise ValueError('Source messages/labels must be literal cells, not formulas.')
            if any(not isinstance(cell.value, str) or not cell.value.strip() for cell in selected):
                raise ValueError('Missing or non-text source value at row ' + str(number))
            participant, identifier, conversation, text, intent = [cell.value for cell in selected]
            if intent.strip() not in INTENTS or len(text) > 2000:
                raise ValueError('Invalid intent or message length at row ' + str(number))
            records.append({'participant_id': participant.strip(), 'id': identifier.strip(),
                            'conversation_id': conversation.strip(), 'text': text,
                            'provided_intent': intent.strip(), 'source_sheet': 'Messages', 'source_row': number})
        return records
    finally:
        workbook.close()


def audit(records):
    if not records or len({row['id'] for row in records}) != len(records):
        raise ValueError('Need nonempty records with unique message IDs.')
    parent = list(range(len(records)))

    def find(index):
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def join(a, b):
        a, b = find(a), find(b)
        if a != b:
            parent[max(a, b)] = min(a, b)

    groups = defaultdict(list)
    seen = {}
    for index, row in enumerate(records):
        key = text_key(row['text'])
        if not key:
            raise ValueError('Empty normalized message.')
        groups[key].append(row)
        # Connected components preserve all three constraints simultaneously.
        for field, value in (('participant', row['participant_id']),
                             ('conversation', row['conversation_id']), ('exact_text', key)):
            group = (field, value)
            if group in seen:
                join(index, seen[group])
            else:
                seen[group] = index
    components = defaultdict(list)
    for index, row in enumerate(records):
        components[find(index)].append(row)
    conflicts, queue, manifest = [], [], []
    for key, rows in sorted(groups.items(), key=lambda pair: pair[1][0]['id']):
        identifier = 'text_' + hashlib.sha256(key.encode()).hexdigest()[:16]
        labels = sorted({row['provided_intent'] for row in rows})
        if len(labels) > 1:
            conflicts.append({'unique_id': identifier, 'provided_intents': labels})
        manifest.append({'unique_id': identifier, 'representative_id': rows[0]['id'],
                         'normalized_text_sha256': hashlib.sha256(key.encode()).hexdigest(),
                         'source_ids': [row['id'] for row in rows], 'occurrence_count': len(rows)})
        queue.append({'id': identifier, 'text': rows[0]['text'], 'provided_intents': labels,
                      'source_ids': [row['id'] for row in rows],
                      'participant_ids': sorted({row['participant_id'] for row in rows}),
                      'conversation_ids': sorted({row['conversation_id'] for row in rows}),
                      'reviewed_intent': None, 'entities': None, 'slots': None, 'remove_slots': None,
                      'reviewer_ids': [], 'family_id': None, 'split': None,
                      'annotation_status': 'pending',
                      'note': 'Read all source conversation contexts before accepting one shared label.'})
    unique_per_label = {label: sum(any(row['provided_intent'] == label for row in rows) for rows in groups.values())
                        for label in INTENTS}
    near_duplicates = []
    for left, first in enumerate(queue):
        for second in queue[left + 1:]:
            ratio = SequenceMatcher(None, text_key(first['text']), text_key(second['text']), autojunk=False).ratio()
            if ratio >= .88:
                near_duplicates.append({'first_id': first['id'], 'second_id': second['id'],
                                        'character_similarity': round(ratio, 4), 'human_family_review': 'pending'})
    report = {'rows': len(records), 'participants': len({row['participant_id'] for row in records}),
              'conversations': len({row['conversation_id'] for row in records}),
              'provided_intent_counts': dict(Counter(row['provided_intent'] for row in records)),
              'unique_texts': len(groups), 'repeated_extra_rows': len(records) - len(groups),
              'unique_texts_per_intent': unique_per_label, 'conflicting_exact_text_labels': conflicts,
              'indivisible_group_count': len(components),
              'groups': [{'group_id': f'group_{index:03}', 'rows': len(rows),
                          'participants': sorted({row['participant_id'] for row in rows}),
                          'source_ids': [row['id'] for row in rows]}
                         for index, rows in enumerate(components.values(), 1)],
              'split_status': 'blocked' if len(components) < 3 else 'requires_family_review_and_class_coverage_check',
              'split_reason': 'Keep participants, conversations and identical texts in a single partition; fewer than three connected groups cannot form train/dev/test.'
                              if len(components) < 3 else 'No splits assigned automatically before paraphrase-family review.',
              'near_duplicate_candidate_pairs': len(near_duplicates),
              'near_duplicate_method': 'Character SequenceMatcher ratio >= 0.88; review aid, not semantic duplicate ground truth.',
              'review_status': 'Provided intents are not independently adjudicated; entity/slot labels remain null.',
              'model_trained': False, 'final_test_evaluated': False,
              'original_rows_deleted': 0}
    return report, queue, manifest, near_duplicates


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--user-confirmed-volunteer-permission', action='store_true', required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise ValueError('Use a new intake output directory; no overwrite allowed.')
    source_hash = hashlib.sha256(args.input.read_bytes()).hexdigest()
    records = read_records(args.input)
    if hashlib.sha256(args.input.read_bytes()).hexdigest() != source_hash:
        raise ValueError('Source changed during reading; rerun after saving it.')
    report, queue, manifest, near = audit(records)
    report.update(source_sha256=source_hash, recorded_utc=datetime.now(timezone.utc).isoformat(),
                  source_filename=args.input.name,
                  provenance='Volunteer-written and permitted for university research, confirmed by the user in chat; consent documents not independently verified.')
    args.output.mkdir(parents=True, exist_ok=False)
    shutil.copyfile(args.input, args.output / 'source_workbook.xlsx')
    if hashlib.sha256((args.output / 'source_workbook.xlsx').read_bytes()).hexdigest() != source_hash:
        raise ValueError('Snapshot identity mismatch; do not use this incomplete intake.')
    for name, value in (('audit_report.json', report), ('source_records.json', records),
                        ('annotation_review_queue.json', queue), ('duplicate_manifest.json', manifest),
                        ('near_duplicate_candidates.json', near)):
        (args.output / name).write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({key: report[key] for key in ('rows', 'unique_texts', 'repeated_extra_rows',
                     'unique_texts_per_intent', 'indivisible_group_count', 'split_status', 'model_trained')}, indent=2))
    print('Intake saved; original unchanged. Review queue is not approved training data:', args.output.resolve())


if __name__ == '__main__': main()
